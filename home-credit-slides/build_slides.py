#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Home Credit 信用違約風險 — 簡報內容（單一資料來源）

此檔案同時產生：
  1) home_credit_slides.pptx      — 可立即檢視的本機簡報
  2) batch_update.json            — 供 `gws slides presentations batchUpdate` 推送到 Google Slides 的請求內容

每張投影片以 dict 描述：
  layout: "TITLE"（封面：CENTERED_TITLE + SUBTITLE）或 "TITLE_AND_BODY"
  title : 標題
  body  : 字串（封面副標）或 list[str]（內文條列）
"""
import json, os

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
DECK_TITLE = "Home Credit 信用風險評估 — AI 輔助授信"

SLIDES = [
    {
        "layout": "TITLE",
        "title": "信用風險評估 — AI 輔助授信",
        "body": "競賽：Kaggle —「Home Credit - Credit Risk Model Stability」(2024 年版)\n從來源資料到風險分數：資料為何能預測、如何前處理與訓練，以及授信人員的實務應用",
    },
    {
        "layout": "TITLE_AND_BODY",
        "title": "競賽版本差異：2024 版 vs 2018 版",
        "body": [
            "本專案採用版本：2024「Home Credit - Credit Risk Model Stability」；2018 舊版為「Home Credit Default Risk」。",
            "來源資料：2024 採多檔分層（主表 base + 各 depth 特徵表）、主鍵 case_id、並含時間欄 WEEK_NUM；2018 以 SK_ID_CURR 為主鍵、bureau / previous_application 等多張 csv。",
            "欄位：2024 欄位以『_型別』後綴匿名命名（_A 金額、_D 日期、_M 類別、_P 逾期…），共數百欄；2018 為具語意名稱（如 AMT_INCOME_TOTAL）。",
            "模型目標（可靠度）：2024 強調『穩定度/可靠度』——不只看單次 AUC，更要求預測表現隨時間（各週）維持穩定，採自訂 gini 穩定度指標，效能下滑會被懲罰；2018 僅以 AUC-ROC 評分。",
            "意涵：2024 版更貼近實務——模型上線後須長期穩定可靠，而非只在歷史測試集上分數高。",
        ],
    },
    {
        "layout": "TITLE_AND_BODY",
        "title": "為什麼這些資料能預測違約風險",
        "body": [
            "資料來源：Kaggle「Home Credit - Credit Risk Model Stability」(2024 年版)——以真實貸款申請人的歷史紀錄，搭配是否違約的結果(target)為標籤，供模型學習。",
            "過往還款行為：逾期天數、準時繳款比例 —— 最直接反映償債意願與能力。",
            "負債與收入結構：授信金額/收入、月攤還/收入比 —— 衡量還款壓力是否過重。",
            "既有信貸狀況：在各機構的貸款筆數與未償餘額 —— 反映整體負債水位。",
            "申請與查詢頻率：短期內多次申請 —— 常是資金緊張的警訊。",
            "金流穩定度與申請人輪廓：收入/帳戶進出穩定度、年齡、年資等 —— 反映財務健康並區隔風險族群。",
            "小結：上述『行為 + 結構』訊號與歷史違約結果高度相關，故模型能據此在統計上區分高/低風險。",
        ],
    },
    {
        "layout": "TITLE_AND_BODY",
        "title": "資料前處理與機器學習流程（總覽）",
        "body": [
            "探索分析(EDA)：先看資料分布、缺失與異常，並注意違約屬於少數族群（不平衡）。",
            "資料清理：修正異常值與日期格式、處理缺失值，讓資料一致可用。",
            "特徵工程：把多筆歷史紀錄彙整成每位申請人的『行為摘要與比率指標』，並挑出最有預測力的特徵。",
            "資料切分(Split)：分訓練/驗證，且依時間切分，確保模型對『未來新案件』仍然準確。",
            "模型訓練與評估：讓模型從歷史案件學習違約樣態，再用沒看過的資料檢驗可靠度與穩定度。",
            "產出：為每一筆申請輸出一個『風險分數』。",
        ],
    },
    {
        "layout": "TITLE_AND_BODY",
        "title": "模型產出：風險分數的意義",
        "body": [
            "模型把眾多訊號濃縮成單一『風險分數』，便於快速判讀。",
            "分數對應違約機率，可切分為 低 / 中 / 高 三個風險等級。",
            "同時列出『主要風險因子』，讓判斷可解釋、可向客戶與主管說明。",
            "重視穩定度：分數須隨時間維持可靠，不因近期樣態變動而失準。",
        ],
    },
    {
        "layout": "TITLE_AND_BODY",
        "title": "實務應用：AI 輔助授信流程",
        "body": [
            "收件 → 帶入/輸入申請人資料 → AI 模型即時產生風險分數與等級。",
            "低風險：快速核准；中風險：人工複審並補件；高風險：加強審查或婉拒。",
            "AI 是『輔助』而非取代：最終決策由授信人員結合分數、因子與經驗判斷。",
            "效益：加速審件、統一風險判讀標準、降低人為遺漏、保留決策軌跡。",
        ],
    },
    {
        "layout": "TITLE_AND_BODY",
        "title": "網頁設計：授信風險評估輔助工具",
        "body": [
            "輸入區：申請人關鍵欄位（年齡、年收入、授信金額、月攤還、年資、過往逾期、既有貸款數、外部信用分數…）。",
            "一鍵『AI 風險評估』，即時計算。",
            "輸出區：風險分數儀表、風險等級、預估違約機率、主要風險因子清單、建議處置。",
            "合規提示：分數為決策輔助，仍須人工複核並符合法遵要求。",
            "（實際互動畫面見隨附網頁 credit_risk_app/index.html）",
        ],
    },
    {
        "layout": "TITLE_AND_BODY",
        "title": "價值與結論",
        "body": [
            "來源資料涵蓋行為與結構訊號 —— 具備足夠的違約預測力。",
            "標準化的前處理與訓練流程 —— 結果穩定、可重現、可解釋。",
            "落地為授信人員的網頁工具 —— 即時風險分數，授信決策更快也更一致。",
            "AI 定位為輔助決策，兼顧效率、風險控管與可解釋性。",
        ],
    },
]


# ---------------------------------------------------------------------------
# 1) 產生 .pptx（本機可立即檢視）
# ---------------------------------------------------------------------------
def build_pptx(path):
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation()
    title_layout = prs.slide_layouts[0]   # Title Slide
    body_layout = prs.slide_layouts[1]    # Title and Content
    for s in SLIDES:
        if s["layout"] == "TITLE":
            slide = prs.slides.add_slide(title_layout)
            slide.shapes.title.text = s["title"]
            if slide.placeholders and len(slide.placeholders) > 1:
                slide.placeholders[1].text = s["body"]
        else:
            slide = prs.slides.add_slide(body_layout)
            slide.shapes.title.text = s["title"]
            tf = slide.placeholders[1].text_frame
            tf.clear()
            for i, line in enumerate(s["body"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = line
                p.font.size = Pt(16)
    prs.save(path)
    print("wrote", path, "(", len(SLIDES), "slides )")


# ---------------------------------------------------------------------------
# 2) 產生 Google Slides batchUpdate 請求（供 gws 推送）
#    使用 createSlide + placeholderIdMappings 取得可控的 objectId，
#    再以 insertText 寫入標題/內文，body 套用項目符號。
# ---------------------------------------------------------------------------
def build_batch_requests():
    requests = []
    for idx, s in enumerate(SLIDES):
        sid = f"slide_{idx}"
        if s["layout"] == "TITLE":
            title_id, sub_id = f"{sid}_title", f"{sid}_sub"
            requests.append({"createSlide": {
                "objectId": sid, "insertionIndex": idx,
                "slideLayoutReference": {"predefinedLayout": "TITLE"},
                "placeholderIdMappings": [
                    {"layoutPlaceholder": {"type": "CENTERED_TITLE", "index": 0}, "objectId": title_id},
                    {"layoutPlaceholder": {"type": "SUBTITLE", "index": 0}, "objectId": sub_id},
                ]}})
            requests.append({"insertText": {"objectId": title_id, "text": s["title"]}})
            requests.append({"insertText": {"objectId": sub_id, "text": s["body"]}})
        else:
            title_id, body_id = f"{sid}_title", f"{sid}_body"
            requests.append({"createSlide": {
                "objectId": sid, "insertionIndex": idx,
                "slideLayoutReference": {"predefinedLayout": "TITLE_AND_BODY"},
                "placeholderIdMappings": [
                    {"layoutPlaceholder": {"type": "TITLE", "index": 0}, "objectId": title_id},
                    {"layoutPlaceholder": {"type": "BODY", "index": 0}, "objectId": body_id},
                ]}})
            requests.append({"insertText": {"objectId": title_id, "text": s["title"]}})
            body_text = "\n".join(s["body"])
            requests.append({"insertText": {"objectId": body_id, "text": body_text}})
            requests.append({"createParagraphBullets": {
                "objectId": body_id,
                "textRange": {"type": "ALL"},
                "bulletPreset": "BULLET_DISC_CIRCLE_SQUARE"}})
    # 移除建立簡報時預設的第一張空白投影片（其 id 於執行期由 create_slides.sh 注入）
    requests.append({"deleteObject": {"objectId": "__DEFAULT_SLIDE_ID__"}})
    return {"requests": requests}


if __name__ == "__main__":
    build_pptx(os.path.join(OUT_DIR, "home_credit_slides.pptx"))
    with open(os.path.join(OUT_DIR, "batch_update.json"), "w", encoding="utf-8") as f:
        json.dump(build_batch_requests(), f, ensure_ascii=False, indent=2)
    print("wrote batch_update.json")
    with open(os.path.join(OUT_DIR, "deck_title.txt"), "w", encoding="utf-8") as f:
        f.write(DECK_TITLE)
